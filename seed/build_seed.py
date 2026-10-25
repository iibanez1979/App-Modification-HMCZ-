"""Convierte el Gantt y los SOP (Excel) en los CSV que usa la app.

Se ejecuta solo cuando cambian los Excel de seed/:

    python seed/build_seed.py

Genera (en seed/):
  activities_seed.csv  Plan original por camión y actividad (desde el Gantt)
  tasks_catalog.csv    Lista normalizada de tareas de los 5 SOP (EN/JP)
  sop_notes.csv        Hojas "Engineering Notes" de cada SOP
"""

from __future__ import annotations

import re
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from core import config  # noqa: E402

warnings.filterwarnings("ignore", module="openpyxl")
JP = re.compile(r"[぀-ヿ一-鿿]")


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in {"#REF!", "None", "nan"} else s


def _date(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return ""


# ─── Gantt ────────────────────────────────────────────────────────────────
def build_activities() -> pd.DataFrame:
    wb = openpyxl.load_workbook(config.SEED / config.GANTT_FILE, data_only=True)
    ws = wb.worksheets[0]
    by_name = {a.gantt_name.lower(): a for a in config.ACTIVITIES}
    rows, truck = [], None
    for r in range(1, ws.max_row + 1):
        item = ws.cell(r, 2).value
        if isinstance(item, str):
            m = re.match(r"RD\s*(\d+)\s*:", item)
            if m:
                truck = f"RD{m.group(1)}"
            continue
        name = _clean(ws.cell(r, 3).value)
        if truck is None or not isinstance(item, (int, float)) or not name:
            continue
        act = by_name.get(name.lower())
        if act is None:
            raise ValueError(f"Fila {r}: actividad desconocida «{name}»")
        progress = ws.cell(r, 5).value
        rows.append({
            "truck": truck,
            "act_key": act.key,
            "base_start": _date(ws.cell(r, 6).value),
            "base_end": _date(ws.cell(r, 7).value),
            "assigned_to": _clean(ws.cell(r, 4).value),
            "imported_progress": float(progress) if isinstance(progress, (int, float)) else 0.0,
        })
    return pd.DataFrame(rows)


# ─── SOP ──────────────────────────────────────────────────────────────────
def _find(headers: list[str], *candidates: str) -> int | None:
    low = [h.lower() for h in headers]
    for cand in candidates:
        for i, h in enumerate(low):
            if h.startswith(cand.lower()):
                return i
    return None


def _qc_gate(raw: str) -> tuple[str, str]:
    """Devuelve (puerta, nota). HOLD = punto de retención; CHECK = verificación."""
    up = raw.upper()
    if up in {"", "NO"}:
        return "", ""
    if up == "YES" or "HOLD" in up:
        return "HOLD", "" if up in {"YES", "HOLD"} else raw
    return "CHECK", "" if up == "CHECK" else raw


def build_catalog() -> tuple[pd.DataFrame, pd.DataFrame]:
    tasks, notes = [], []
    for sop, filename in config.SOP_FILES.items():
        wb = openpyxl.load_workbook(config.SEED / filename, data_only=True)
        ws = wb.worksheets[0]
        headers = [_clean(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
        col = {
            "task_id": _find(headers, "Task ID"),
            "stage": _find(headers, "Stage", "Phase"),
            "desc_en": _find(headers, "Task Description (EN)", "Activity Description (EN)", "Task / Field Instruction"),
            "desc_jp": _find(headers, "作業内容", "作業名"),
            "instr_en": _find(headers, "Detailed Field Instruction (EN)", "Field Instruction (EN)"),
            "instr_jp": _find(headers, "現場作業説明"),
            "drawing": _find(headers, "Drawing / Detail", "Manual Page", "Drawing Page", "Drawing Reference", "Source / Notes"),
            "location": _find(headers, "Location"),
            "predecessor": _find(headers, "Predecessor"),
            "planned_mh": _find(headers, "Planned MH"),
            "qc": _find(headers, "QC Hold Point", "QC Check", "QC Gate", "Hold Point"),
            "remarks": _find(headers, "Remarks"),
        }
        seq = 0
        for r in range(2, ws.max_row + 1):
            def get(key: str) -> str:
                i = col[key]
                return _clean(ws.cell(r, i + 1).value) if i is not None else ""

            task_id = re.sub(r"\s+", "", get("task_id"))
            if not task_id:
                continue
            seq += 1
            stage, desc_en, remarks = get("stage"), get("desc_en"), get("remarks")
            conditional = (
                "conditional" in stage.lower()
                or "n/a if" in remarks.lower()
                or bool(re.search(r"\((conditional|if present)", desc_en.lower()))
            )
            gate, qc_note = _qc_gate(get("qc"))
            drawing = get("drawing")
            if get("location"):
                drawing = f"{get('location')} · {drawing}" if drawing else get("location")
            mh = get("planned_mh")
            tasks.append({
                "sop": sop, "seq": seq, "task_id": task_id, "stage": stage,
                "desc_en": desc_en, "desc_jp": get("desc_jp"),
                "instr_en": get("instr_en"), "instr_jp": get("instr_jp"),
                "drawing": drawing, "predecessor": get("predecessor"),
                "planned_mh": float(mh) if re.fullmatch(r"\d+(\.\d+)?", mh) and float(mh) > 0 else None,
                "qc_gate": gate, "qc_note": qc_note, "conditional": conditional,
            })

        # Hoja "Engineering Notes"
        if len(wb.worksheets) > 1:
            ns = wb.worksheets[1]
            for r in range(2, ns.max_row + 1):
                vals = [_clean(ns.cell(r, c).value) for c in range(1, ns.max_column + 1)]
                vals = [v for v in vals]
                if not vals or not vals[0]:
                    continue
                rest = [v for v in vals[1:] if v]
                en = next((v for v in rest if not JP.search(v)), "")
                jp = next((v for v in rest if JP.search(v)), "")
                source = rest[-1] if len(rest) >= 3 and rest[-1] not in {en, jp} else ""
                notes.append({"sop": sop, "item": vals[0], "note_en": en, "note_jp": jp, "source": source})
    return pd.DataFrame(tasks), pd.DataFrame(notes)


def main() -> None:
    acts = build_activities()
    acts.to_csv(config.ACTIVITIES_SEED, index=False)
    cat, notes = build_catalog()
    cat.to_csv(config.TASKS_CATALOG, index=False)
    notes.to_csv(config.SOP_NOTES, index=False)
    print(f"activities_seed.csv: {len(acts)} filas · {acts.truck.nunique()} camiones")
    print(f"tasks_catalog.csv:   {len(cat)} tareas · {cat.conditional.sum()} condicionales")
    print(cat.groupby("sop").size().to_string())
    print(f"sop_notes.csv:       {len(notes)} notas")


if __name__ == "__main__":
    main()
