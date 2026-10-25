"""Pruebas de la app completa con AppTest (modo local en una carpeta temporal)."""

from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

APP = str(Path(__file__).resolve().parent.parent / "app.py")


@pytest.fixture
def at(tmp_path, monkeypatch):
    monkeypatch.setenv("HMCZ_LOCAL_DIR", str(tmp_path / "data"))
    import streamlit as st
    st.cache_resource.clear()
    st.cache_data.clear()
    return AppTest.from_file(APP, default_timeout=60)


def _setup_admin(at):
    at.run()
    at.text_input[0].input("ivan")
    at.text_input[1].input("Ivan Ibanez")
    at.text_input[2].input("secret123")
    at.text_input[3].input("secret123")
    at.button[0].click().run()
    return at


def test_first_admin_then_admin_view(at):
    _setup_admin(at)
    assert not at.exception
    assert len(at.tabs) == 7
    assert any("Coordinate trunnion work" in m.value for m in at.markdown)


def test_login_fails_with_wrong_password(at, tmp_path):
    _setup_admin(at)
    at.session_state["user"]  # sesión iniciada
    at2 = AppTest.from_file(APP, default_timeout=60).run()
    at2.text_input[0].input("ivan")
    at2.text_input[1].input("wrong-password")
    at2.button[0].click().run()
    assert "user" not in at2.session_state
    assert at2.error


def test_supervisor_view_renders(at):
    _setup_admin(at)
    at.session_state["view_pick"] = "supervisor"
    at.run()
    assert not at.exception
    assert at.selectbox[0].options[0].startswith("RD")


def test_japanese_ui(at):
    _setup_admin(at)
    at.session_state["lang_pick"] = "ja"
    at.run()
    assert not at.exception
    assert any("トラニオン" in m.value for m in at.markdown)


def test_admin_actions_shift_ack_settings(at, tmp_path):
    _setup_admin(at)
    truck = at.selectbox(key="plan_truck").value
    at.text_input(key=f"shift_reason_{truck}").input("Parts late")
    next(b for b in at.button if b.label == "Shift and save").click().run()
    assert not at.exception
    assert list((tmp_path / "data" / "log" / "plan").rglob("*.json"))
    at.text_input(key="ack_RD224").input("PO 123")
    at.button(key="ackb_RD224").click().run()
    assert list((tmp_path / "data" / "log" / "ack").rglob("*.json"))
    assert not any("RD224 · Coordinate" in m.value for m in at.markdown)


def test_supervisor_delay_note(at, tmp_path):
    _setup_admin(at)
    at.session_state["view_pick"] = "supervisor"
    at.run()
    area = next(a for a in at.text_area if a.label == "Detailed description")
    area.input("Waiting for cover plates")
    next(b for b in at.button if b.label == "Submit delay note").click().run()
    assert not at.exception
    assert list((tmp_path / "data" / "log" / "delay").rglob("*.json"))


def test_excel_report(tmp_path):
    from datetime import date
    from openpyxl import load_workbook
    from core.events import initialize
    from core.project import build_project
    from core.reports import build_excel
    from storage.local_store import LocalStore

    s = LocalStore(tmp_path / "d")
    initialize(s, "x")
    data = build_excel(build_project(s.snapshot(), today=date(2026, 10, 25)))
    import io
    wb = load_workbook(io.BytesIO(data))
    assert {"Summary", "Trucks", "Activities", "SOP tasks", "Delay notes"} <= set(wb.sheetnames)
    assert wb["Trucks"].max_row == 13
